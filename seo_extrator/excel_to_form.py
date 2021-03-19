from selenium import webdriver
from selenium.webdriver.support.ui import WebDriverWait
import pdb
import sys
from selenium.webdriver.support import expected_conditions as EC
from openpyxl import load_workbook
import docx
import math

wb = load_workbook(filename = 'inputs.xlsm')
sheet_ranges = wb['Planilha1']
sheet_ranges2 = wb['Planilha2']


#contar quantas linas preenchidas há na planilha
ws = wb.active
celulas_preenchidas = tuple(ws.rows)
contar = int(0)
for x in celulas_preenchidas:
    contar += 1

#vamos extrair os textos
    
document = docx.Document("carta1.docx")
p = document.paragraphs
tamanho_txt = len(p)
y=""
for x in range(tamanho_txt):
 y += "\n" + p[x].text


coments = y.split('#$%')



driver = webdriver.Chrome()

driver.get(sheet_ranges['a4'].value)



try:
    driver.find_element_by_css_selector('textarea').send_keys(coments[0])
except:
    print('Não foi localizado textarea nesse site')



for x in range(3, contar + 1):


    

    num = x - 3

    site = "d{}".format(x)
    email3 = "b{}".format(x)
    assunto = "c{}".format(x)
    titulo3 = "e{}".format(x)
  

    driver.implicitly_wait(10)
    b = str(sheet_ranges[site].value)
    driver.execute_script("window.open('"+ b + "', '" + "new_window{}".format(num) +"')")
    WebDriverWait(driver, 10).until(EC.number_of_windows_to_be(x-1))
    window_after = driver.window_handles[num + 1]
    driver.switch_to.window(window_after)

    try:
        driver.find_element_by_css_selector('textarea').send_keys(coments[x - 2])
    except:
        print("Não foi localizado um campo para colocar o texto em " + b )

    
    try:         
        driver.find_element_by_css_selector("["+sheet_ranges2[email3].value+"]").send_keys("patiocontins@estacionamentopatioconfins.com.br")
    except:
        print("Não foi localizado um campo para colocar o e-mail em " + b )


    try:        
        driver.find_element_by_css_selector("["+sheet_ranges2[assunto].value+"]").send_keys(sheet_ranges[titulo3].value)
    except:
        print("Não foi localizado um campo para colocar o título em " + b )

    
    if math.fmod(x,12) ==0:
        reply = str(input('Deu tudo  certo? (y/n):')).lower().strip()
        if reply[0] == 'y':

            pass
        elif reply[0] == 'n':
            sys.exit()
        
        else:
            print("prosseguindo...")

    if x == contar +1:
        input('Pode fechar?')




#driver.close()
