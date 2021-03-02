from selenium import webdriver
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
import pdb
import time
from selenium.webdriver.support import expected_conditions as EC
from openpyxl import load_workbook
import docx

wb = load_workbook(filename = 'ListadeEmails.xlsx')
sheet_ranges = wb['Planilha1']



#contar quantas linas preenchidas há na planilha
ws = wb.active
celulas_preenchidas = tuple(ws.rows)
contar = int(0)
for x in celulas_preenchidas:
    contar += 1

#vamos extrair os textos
    
document = docx.Document("1.docx")
p = document.paragraphs
tamanho_txt = len(p)
y=""
for x in range(tamanho_txt):
 y += "\n" + p[x].text


coments = y.split('#$%')

#senh = input("Coloque a senha")
senh = "pconf209"

driver = webdriver.Chrome()

driver.get("https://email.uolhost.com.br/")

WebDriverWait(driver, timeout=3).until(lambda d: d.find_element(By.CSS_SELECTOR, "input[type='password']"))

emailr = driver.find_element(By.CSS_SELECTOR, "input[type='email']")

emailr.send_keys("patioconfins@estacionamentopatioconfins.com.br")

senha = driver.find_element(By.CSS_SELECTOR, "input[type='password']")

senha.send_keys(senh)

driver.find_element(By.CSS_SELECTOR, "#login > div > form > button").click()

for x in range(2, contar + 1):

    endereco = "d{}".format(x)
    titulo_msg = "f{}".format(x)
    

    escrever = WebDriverWait(driver, timeout=20).until(lambda e: e.find_element(By.CSS_SELECTOR, "menu[ng-click='onComposeNew($event)']"))

    try:
        escrever.click()
    except:
        time.sleep(4)
        escrever.click()

    iframe = WebDriverWait(driver, timeout=30).until(lambda f: f.find_element(By.CSS_SELECTOR, "iframe[aria-describedby='cke_40']"))



    #email

    endereco_email = driver.find_element(By.CSS_SELECTOR, "input[realfieldid='field-to']")
    endereco_email.send_keys(sheet_ranges[endereco].value)

    #assunto
    assunto = driver.find_element(By.CSS_SELECTOR, "input[ng-model='subject']")
    assunto.send_keys(sheet_ranges[titulo_msg].value)
    time.sleep(1)
    assunto.send_keys(Keys.TAB + Keys.TAB)
    time.sleep(1)

    #texto  worContentEditable_field-body
    driver.switch_to.active_element.send_keys(coments[x-2])


    #enviar
    try:
        
        driver.find_element(By.CSS_SELECTOR, "menu[ng-click='onSend()']").click()
    except:
        time.sleep(4)
        driver.find_element(By.CSS_SELECTOR, "body > div.flex.flex-box-v.viewport > div > div.view-modal.webmail_accounts_folders_compose > form > div.toolbar > div.group-bt-left > menu.bt-compose-send.highlight").click()

    finally:
        pdb.set_trace()

    WebDriverWait(driver, timeout=10).until(EC.visibility_of(By.ID, "notifications"))











    




driver.close()
